"""
STEP 1 — Document loading.

Har file ko `Document(text, metadata)` mein badalte hain.
Metadata mein `source` (file ka naam) aur `page` (PDF ka page number) rakhte hain,
taaki baad mein answer ke saath citation de sakein: "(notes.pdf, page 4)".
"""
from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, List

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".csv", ".html", ".htm"}


@dataclass
class Document:
    text: str
    metadata: dict = field(default_factory=dict)


def file_hash(path: Path) -> str:
    """File ka content hash — isse pata chalta hai file badli hai ya nahi (re-ingest skip karne ke liye)."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)          # extra spaces
    text = re.sub(r"\n{3,}", "\n\n", text)        # 3+ newlines -> paragraph break
    return text.strip()


def _load_pdf(path: Path) -> List[Document]:
    # PyMuPDF (fitz) zyada achha text nikalta hai; na ho to pypdf
    docs: List[Document] = []
    try:
        import fitz  # PyMuPDF

        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf):
                text = _clean(page.get_text())
                if text:
                    docs.append(Document(text, {"page": i + 1}))
        return docs
    except ImportError:
        pass

    from pypdf import PdfReader

    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages):
        text = _clean(page.extract_text() or "")
        if text:
            docs.append(Document(text, {"page": i + 1}))
    return docs


def _load_text(path: Path) -> List[Document]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [Document(_clean(text), {})]


def _load_docx(path: Path) -> List[Document]:
    import docx  # python-docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Document(_clean("\n".join(parts)), {})]


def _load_csv(path: Path) -> List[Document]:
    """Har row ek alag Document — 'column: value' format mein, taaki LLM samajh sake."""
    docs: List[Document] = []
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader):
            text = "\n".join(f"{k}: {v}" for k, v in row.items() if k is not None)
            if text.strip():
                docs.append(Document(text, {"row": i + 1}))
    return docs


def _load_html(path: Path) -> List[Document]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()
        text = soup.get_text("\n")
    except ImportError:
        text = re.sub(r"<script.*?</script>|<style.*?</style>", " ", raw, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    return [Document(_clean(text), {})]


_LOADERS = {
    ".pdf": _load_pdf,
    ".txt": _load_text,
    ".md": _load_text,
    ".docx": _load_docx,
    ".csv": _load_csv,
    ".html": _load_html,
    ".htm": _load_html,
}


def load_file(path: Path) -> List[Document]:
    path = Path(path)
    loader = _LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    docs = loader(path)
    for d in docs:
        d.metadata["source"] = path.name
        d.metadata["path"] = str(path)
    return [d for d in docs if d.text]


def iter_files(root: Path) -> Iterable[Path]:
    root = Path(root)
    if root.is_file():
        if root.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield root
        return
    for p in sorted(root.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS and not p.name.startswith("."):
            yield p
